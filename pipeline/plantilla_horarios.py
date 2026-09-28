#!/usr/bin/env python3
"""
Genera la plantilla que AEMUS debe completar con la informacion operativa.

Diseno: una fila por BLOQUE HORARIO, no columnas fijas por franja. Cada ruta
puede tener tantos bloques como necesite (dos picos de manana, un pico largo,
una frecuencia distinta a media tarde), que es como opera de verdad y ademas
como el GTFS guarda las frecuencias: desde / hasta / cada cuantos minutos.

Salida: PLANTILLA_HORARIOS_AEMUS.xlsx (3 hojas) y su equivalente en CSV.
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from etl import AGENCIAS, RUTAS  # noqa: E402

AQUI = os.path.dirname(os.path.abspath(__file__))

AMARILLO = "FFD000"
OSCURO = "2C3438"
CREMA = "FFF8CC"
GRIS = "F2F6F7"

TIPOS_DIA = ["Lunes a viernes", "Sábado", "Domingo y feriados"]

CAB = ["Código", "Empresa", "Ruta", "Sentido", "Tipo de día",
       "Desde", "Hasta", "Cada cuántos minutos", "Observaciones"]

EJEMPLO = [
    ["9999", "(EJEMPLO — no completar esta fila)", "RUTA DE MUESTRA", "IDA",
     "Lunes a viernes", "05:00", "06:30", "12", "inicio de operación"],
    ["9999", "(EJEMPLO)", "RUTA DE MUESTRA", "IDA", "Lunes a viernes",
     "06:30", "09:00", "5", "hora punta mañana"],
    ["9999", "(EJEMPLO)", "RUTA DE MUESTRA", "IDA", "Lunes a viernes",
     "09:00", "17:00", "10", ""],
    ["9999", "(EJEMPLO)", "RUTA DE MUESTRA", "IDA", "Lunes a viernes",
     "17:00", "20:30", "6", "hora punta tarde"],
    ["9999", "(EJEMPLO)", "RUTA DE MUESTRA", "IDA", "Lunes a viernes",
     "20:30", "22:45", "15", "última salida 22:45"],
]

NOTAS = [
    ("Qué necesitamos", ""),
    ("", "La información operativa de cada ruta: en qué horario opera, cada cuántos "
         "minutos sale una unidad, qué días y cuánto cuesta el pasaje."),
    ("", ""),
    ("Cómo se completa", ""),
    ("", "Una fila por cada bloque horario con frecuencia distinta. Si una ruta "
         "mantiene la misma frecuencia todo el día, basta una sola fila. Si tiene "
         "dos horas punta, se agregan las filas que hagan falta."),
    ("", "Los bloques deben cubrir sin huecos desde la primera hasta la última "
         "salida. El «Hasta» de un bloque es el «Desde» del siguiente."),
    ("", "Puede insertar todas las filas que necesite. La hoja trae una fila por "
         "ruta, sentido y tipo de día solo como punto de partida."),
    ("", ""),
    ("Formato de los datos", ""),
    ("", "Horas en formato de 24 horas: 05:00, 14:30, 22:45."),
    ("", "Frecuencia en minutos, solo el número: 12 (no «cada 12 minutos»)."),
    ("", "Si un tipo de día no tiene servicio, escriba «no opera» en Observaciones "
         "y deje las horas en blanco."),
    ("", ""),
    ("Si no hay frecuencia fija", ""),
    ("", "Si el despacho no se maneja por intervalos sino por una tabla de salidas "
         "con horas exactas, envíenos esa tabla tal como la manejan: nos sirve "
         "igual o mejor. No hace falta convertirla."),
    ("", ""),
    ("Dudas", ""),
    ("", "Con gusto la completamos juntos en la reunión semanal de los martes. "
         "Suele tomar menos tiempo que resolverlo por correo."),
]


def estilo_cabecera(ws, fila, ncols):
    borde = Side(style="thin", color="D7DEE0")
    for c in range(1, ncols + 1):
        cel = ws.cell(row=fila, column=c)
        cel.font = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=OSCURO)
        cel.alignment = Alignment(vertical="center", wrap_text=True)
        cel.border = Border(bottom=borde)
    ws.row_dimensions[fila].height = 30


def main():
    wb = Workbook()

    # ---------- Hoja 1: instrucciones ----------
    ws = wb.active
    ws.title = "Instrucciones"
    ws["A1"] = "Información operativa de rutas — AEMUS"
    ws["A1"].font = Font(name="Calibri", bold=True, size=16, color=OSCURO)
    ws["A2"] = "Trufi Association e.V. · construcción del GTFS de Lima"
    ws["A2"].font = Font(name="Calibri", size=11, color="6E7A80")
    ws.merge_cells("A1:B1")
    ws.merge_cells("A2:B2")

    fila = 4
    for titulo, texto in NOTAS:
        if titulo:
            ws.cell(row=fila, column=1, value=titulo).font = Font(
                name="Calibri", bold=True, size=11, color=OSCURO)
        if texto:
            cel = ws.cell(row=fila, column=2, value=texto)
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            cel.font = Font(name="Calibri", size=10)
            ws.row_dimensions[fila].height = 15 * (1 + len(texto) // 82)
        fila += 1
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 88

    # ---------- Hoja 2: frecuencias ----------
    ws = wb.create_sheet("Frecuencias")
    ws.append(CAB)
    estilo_cabecera(ws, 1, len(CAB))

    borde = Side(style="thin", color="E4EAEC")
    for ej in EJEMPLO:
        ws.append(ej)
        for c in range(1, len(CAB) + 1):
            cel = ws.cell(row=ws.max_row, column=c)
            cel.fill = PatternFill("solid", fgColor=CREMA)
            cel.font = Font(name="Calibri", size=10, italic=True, color="8A6D0B")
            cel.border = Border(bottom=borde)

    ws.append([])
    for ruta in RUTAS:
        empresa = next(a[1] for a in AGENCIAS if a[0] == ruta["agency_id"])
        for sentido in ("IDA", "RETORNO"):
            for tipo in TIPOS_DIA:
                ws.append([ruta["codigo"], empresa, ruta["marca"], sentido,
                           tipo, "", "", "", ""])
                for c in range(1, len(CAB) + 1):
                    cel = ws.cell(row=ws.max_row, column=c)
                    cel.font = Font(name="Calibri", size=10)
                    cel.border = Border(bottom=borde)
                    if c <= 5:
                        cel.fill = PatternFill("solid", fgColor=GRIS)
                    else:
                        cel.fill = PatternFill("solid", fgColor="FFFFFF")

    for col, ancho in zip("ABCDEFGHI", (9, 34, 21, 11, 19, 9, 9, 20, 30)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "F2"

    # ---------- Hoja 3: tarifas ----------
    ws = wb.create_sheet("Tarifas")
    cab2 = ["Código", "Empresa", "Ruta", "Tipo de pasaje",
            "Tarifa (S/)", "Observaciones"]
    ws.append(cab2)
    estilo_cabecera(ws, 1, len(cab2))
    ws.append(["9999", "(EJEMPLO — no completar)", "RUTA DE MUESTRA",
               "Adulto", "2.50", "tarifa plana en toda la ruta"])
    ws.append(["9999", "(EJEMPLO)", "RUTA DE MUESTRA", "Escolar / universitario",
               "1.00", ""])
    for r in (2, 3):
        for c in range(1, len(cab2) + 1):
            cel = ws.cell(row=r, column=c)
            cel.fill = PatternFill("solid", fgColor=CREMA)
            cel.font = Font(name="Calibri", size=10, italic=True, color="8A6D0B")
    ws.append([])
    for ruta in RUTAS:
        empresa = next(a[1] for a in AGENCIAS if a[0] == ruta["agency_id"])
        for tipo in ("Adulto", "Escolar / universitario"):
            ws.append([ruta["codigo"], empresa, ruta["marca"], tipo, "", ""])
            for c in range(1, len(cab2) + 1):
                cel = ws.cell(row=ws.max_row, column=c)
                cel.font = Font(name="Calibri", size=10)
                if c <= 4:
                    cel.fill = PatternFill("solid", fgColor=GRIS)
    for col, ancho in zip("ABCDEF", (9, 34, 21, 24, 13, 34)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "E2"

    destino = os.path.join(AQUI, "PLANTILLA_HORARIOS_AEMUS.xlsx")
    wb.save(destino)

    # ---------- Equivalente en CSV ----------
    csv_destino = os.path.join(AQUI, "PLANTILLA_HORARIOS_AEMUS.csv")
    with open(csv_destino, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(CAB)
        w.writerows(EJEMPLO)
        w.writerow([])
        for ruta in RUTAS:
            empresa = next(a[1] for a in AGENCIAS if a[0] == ruta["agency_id"])
            for sentido in ("IDA", "RETORNO"):
                for tipo in TIPOS_DIA:
                    w.writerow([ruta["codigo"], empresa, ruta["marca"],
                                sentido, tipo, "", "", "", ""])

    print("✔", destino)
    print("✔", csv_destino)


if __name__ == "__main__":
    main()
