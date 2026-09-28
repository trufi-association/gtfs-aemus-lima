#!/usr/bin/env python3
"""Completa el GTFS de AEMUS con las tablas que dependen de los datos operativos.

Entradas
  ../../PLANTILLA_HORARIOS_AEMUS_PRELLENADA.xlsx   frecuencias y tarifas devueltas por AEMUS
  ../../tiempos_recorrido/TIEMPOS_RECORRIDO_AEMUS_<fecha>.csv  tiempos por ruta/sentido/franja
  secuencia_paradas.csv + gtfs/stops.txt + gtfs/shapes.txt    (salida de etl.py)

Salidas (en gtfs/)
  calendar.txt, calendar_dates.txt, trips.txt (reescrito), stop_times.txt,
  frequencies.txt, fare_attributes.txt, fare_rules.txt, feed_info.txt
  y el zip ../gtfs_aemus_lima_<fecha>.zip

Criterios (ver DECISIONES_GTFS_AEMUS.md):
  - Feed basado en frecuencias: un trip por ruta, sentido, servicio y franja de la
    plantilla; stop_times describe el patron desde 00:00:00 y frequencies.txt fija
    las salidas reales.
  - Tiempo de recorrido de cada franja = el de la banda horaria de la tabla de
    tiempos que contiene el punto medio de la franja.
  - Tiempos entre paraderos proporcionales a la distancia recorrida sobre el trazado.
  - Franjas nocturnas (IDA NOCHE / RETORNO NOCHE, 00:00-05:00) se cuelgan del dia
    de servicio anterior como horas 24:00-29:00.
  - "Domingo y feriados": los feriados nacionales del Peru usan el servicio DOM.
  - Tarifa: una por ruta, la maxima de adulto (las tarifas por tramo no caben en
    fare v1); se documenta en fare_attributes.
"""
import csv, datetime as dt, glob, os, sys, zipfile
from collections import defaultdict
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from etl import RUTAS

AQUI = os.path.dirname(os.path.abspath(__file__))
CLI = os.path.abspath(os.path.join(AQUI, "..", ".."))
GTFS = os.path.join(AQUI, "gtfs")
PLANTILLA = os.path.join(CLI, "PLANTILLA_HORARIOS_AEMUS_PRELLENADA.xlsx")
TIEMPOS = sorted(glob.glob(os.path.join(CLI, "tiempos_recorrido", "TIEMPOS_RECORRIDO_AEMUS_*.csv")))[-1]

HOY = dt.date.today()
INICIO, FIN = HOY, HOY + dt.timedelta(days=364)
# Feriados nacionales del Peru dentro de la ventana del feed.
FERIADOS = ["2026-10-08", "2026-11-01", "2026-12-08", "2026-12-09", "2026-12-25",
            "2027-01-01", "2027-03-25", "2027-03-26", "2027-05-01", "2027-06-07",
            "2027-06-29", "2027-07-23", "2027-07-28", "2027-07-29", "2027-08-06", "2027-08-30"]
SERVICIOS = {  # service_id: (lun..dom)
    "LV": (1, 1, 1, 1, 1, 0, 0), "SAB": (0, 0, 0, 0, 0, 1, 0), "DOM": (0, 0, 0, 0, 0, 0, 1),
    "LS": (1, 1, 1, 1, 1, 1, 0), "LD": (1, 1, 1, 1, 1, 1, 1)}
TIPO_DIA = {"lunes a viernes": "LV", "sábado": "SAB", "sabado": "SAB", "domingo y feriados": "DOM",
            "lunes a sabados": "LS", "lunes a sábados": "LS", "lunes a domingo": "LD"}
COD2RID = {r["codigo"]: r["route_id"] for r in RUTAS}


def a_min(v):
    """Celda de hora -> minutos desde 00:00 (acepta time, timedelta o 'HH:MM')."""
    if isinstance(v, dt.time): return v.hour * 60 + v.minute
    if isinstance(v, dt.timedelta): return int(v.total_seconds() // 60)
    h, m = str(v).strip().split(":")[:2]; return int(h) * 60 + int(m)


def hms(mins, seg=0):
    return "%02d:%02d:%02d" % (mins // 60, mins % 60, seg)


def leer_plantilla():
    wb = load_workbook(PLANTILLA, data_only=True)
    franjas = []
    for r in wb["Frecuencias"].iter_rows(min_row=2, values_only=True):
        cod, _, _, sent, tipo, d, h, f, *_ = r
        if not cod or d is None: continue
        cod = str(cod).strip().replace(".0", ""); sent = sent.strip().upper()
        d, h = a_min(d), a_min(h)
        if "NOCHE" in sent: d += 1440; h += 1440          # madrugada del dia siguiente
        elif h <= d: h += 1440                              # cruza medianoche (hasta 00:00)
        franjas.append({"rid": COD2RID[cod], "dir": 0 if sent.startswith("IDA") else 1,
                        "svc": TIPO_DIA[tipo.strip().lower()], "ini": d, "fin": h,
                        "headway": int(float(f))})
    tarifas = defaultdict(float)
    for r in wb["Tarifas"].iter_rows(min_row=2, values_only=True):
        cod, _, _, tipo, precio, *_ = r
        if not cod or precio in (None, ""): continue
        cod = str(cod).strip().replace(".0", "")
        if str(tipo).lower().startswith("adulto"):
            tarifas[COD2RID[cod]] = max(tarifas[COD2RID[cod]], float(precio))
    return franjas, tarifas


def leer_tiempos():
    out = {}
    for r in csv.DictReader(open(TIEMPOS, encoding="utf-8")):
        rid = COD2RID[r["Código"]]; d = 0 if r["Sentido"] == "IDA" else 1
        a, b = r["Franja horaria"].split("-")
        out.setdefault((rid, d), []).append((a_min(a), a_min(b), int(r["Tiempo de recorrido (min)"])))
    return out


def tiempo_para(tiempos, rid, d, ini, fin):
    medio = ((ini + fin) // 2) % 1440
    for a, b, t in tiempos[(rid, d)]:
        if a <= medio < b: return t
    return tiempos[(rid, d)][-1][2]


def leer_paradas():
    """(rid, dir) -> lista de (stop_id, recorrido_m) en orden; empareja secuencia con stops.txt por coordenadas."""
    ids = {}
    for s in csv.DictReader(open(os.path.join(GTFS, "stops.txt"), encoding="utf-8")):
        ids[(s["stop_id"].split("_")[0], s["stop_id"].split("_")[1], s["stop_lat"], s["stop_lon"])] = s["stop_id"]
    out = defaultdict(list)
    for r in csv.DictReader(open(os.path.join(AQUI, "secuencia_paradas.csv"), encoding="utf-8")):
        k = (r["route_id"], r["direction_id"], r["lat"], r["lon"])
        if k in ids: out[(r["route_id"], int(r["direction_id"]))].append((ids[k], r["stop_name"], float(r["recorrido_m"])))
    return out


def largo_shapes():
    L = {}
    for r in csv.DictReader(open(os.path.join(GTFS, "shapes.txt"))):
        L[r["shape_id"]] = max(L.get(r["shape_id"], 0), float(r["shape_dist_traveled"]))
    return L


def escribir(nombre, cab, filas):
    with open(os.path.join(GTFS, nombre), "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh); w.writerow(cab); w.writerows(filas)
    print("  %-20s %5d filas" % (nombre, len(filas)))


def main():
    franjas, tarifas = leer_plantilla()
    tiempos = leer_tiempos(); paradas = leer_paradas(); L = largo_shapes()
    usados = sorted({f["svc"] for f in franjas})
    escribir("calendar.txt", ["service_id", "monday", "tuesday", "wednesday", "thursday", "friday",
                              "saturday", "sunday", "start_date", "end_date"],
             [[s, *SERVICIOS[s], INICIO.strftime("%Y%m%d"), FIN.strftime("%Y%m%d")] for s in usados])
    cd = []
    for f in FERIADOS:
        dia = dt.date.fromisoformat(f)
        if not (INICIO <= dia <= FIN): continue
        ymd = dia.strftime("%Y%m%d"); wd = dia.weekday()
        for s in usados:
            activo = SERVICIOS[s][wd]
            if s == "DOM" and not activo: cd.append([s, ymd, 1])
            elif s in ("LV", "SAB", "LS") and activo: cd.append([s, ymd, 2])
    escribir("calendar_dates.txt", ["service_id", "date", "exception_type"], cd)

    trips, st, fq, n = [], [], [], defaultdict(int)
    for f in sorted(franjas, key=lambda x: (x["rid"], x["dir"], x["svc"], x["ini"])):
        rid, d, svc = f["rid"], f["dir"], f["svc"]
        n[(rid, d, svc)] += 1
        tid = "%s_%d_%s_%02d" % (rid, d, svc, n[(rid, d, svc)])
        pts = paradas[(rid, d)]
        T = tiempo_para(tiempos, rid, d, f["ini"], f["fin"])          # min terminal a terminal
        r0 = pts[0][2]; tramo = pts[-1][2] - r0
        dur = T * tramo / L["%s_%d" % (rid, d)]                       # min entre 1er y ultimo paradero
        trips.append([tid, rid, svc, pts[-1][1], d, "%s_%d" % (rid, d)])
        rec_prev = None
        for i, (sid, _, rec) in enumerate(pts, 1):
            if rec_prev is not None and rec <= rec_prev:
                rec = rec_prev + 1.0   # paraderos que proyectan al mismo punto (fin de trazado): +1 m para mantener orden estricto
            rec_prev = rec
            seg = int(round(dur * 60 * (rec - r0) / tramo)) if tramo else 0
            t = hms(seg // 60, seg % 60)
            st.append([tid, t, t, sid, i, "%.0f" % rec, 1 if i in (1, len(pts)) else 0])
        fq.append([tid, hms(f["ini"]), hms(f["fin"]), f["headway"] * 60, 0])
    escribir("trips.txt", ["trip_id", "route_id", "service_id", "trip_headsign", "direction_id", "shape_id"], trips)
    escribir("stop_times.txt", ["trip_id", "arrival_time", "departure_time", "stop_id", "stop_sequence",
                                "shape_dist_traveled", "timepoint"], st)
    escribir("frequencies.txt", ["trip_id", "start_time", "end_time", "headway_secs", "exact_times"], fq)
    escribir("fare_attributes.txt", ["fare_id", "price", "currency_type", "payment_method", "transfers", "agency_id"],
             [["F_" + rid, "%.2f" % p, "PEN", 0, 0, next(r["agency_id"] for r in RUTAS if r["route_id"] == rid)]
              for rid, p in sorted(tarifas.items())])
    escribir("fare_rules.txt", ["fare_id", "route_id"], [["F_" + rid, rid] for rid in sorted(tarifas)])
    escribir("feed_info.txt", ["feed_publisher_name", "feed_publisher_url", "feed_lang", "feed_start_date",
                               "feed_end_date", "feed_version"],
             [["Trufi Association e.V.", "https://www.trufi-association.org", "es",
               INICIO.strftime("%Y%m%d"), FIN.strftime("%Y%m%d"), "aemus-" + HOY.isoformat()]])
    zpath = os.path.join(AQUI, "..", "gtfs_aemus_lima_%s.zip" % HOY.isoformat())
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for fn in sorted(os.listdir(GTFS)):
            if fn.endswith(".txt"): z.write(os.path.join(GTFS, fn), fn)
    print("✔", os.path.abspath(zpath))


if __name__ == "__main__":
    main()
